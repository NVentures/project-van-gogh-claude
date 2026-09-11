"""deliverable: staging a plan and building the file it describes.

The interesting cases are all about what must not reach a reader. A deck with a
blank slide, a placeholder, or a dash in it is worse than a build that refused,
because the refusal is visible and the bad deck gets sent.
"""
import json

import pytest

import deliverable


# ── Extracting the model's JSON ───────────────────────────────────────────────

def test_a_plan_wrapped_in_prose_is_still_read():
    # `claude -p` prefixes reasoning often enough that a bare json.loads is a
    # latent failure, and the CLI has been seen writing a warning line first.
    reply = 'Here is the deck you asked for:\n{"title": "T", "slides": []}\nHope that helps.'
    assert deliverable.extract_json(reply) == {"title": "T", "slides": []}


def test_a_brace_inside_a_string_does_not_end_the_object():
    reply = '{"title": "a } brace", "slides": []}'
    assert deliverable.extract_json(reply)["title"] == "a } brace"


def test_an_escaped_quote_does_not_end_the_string():
    reply = r'{"title": "she said \"go\"", "slides": []}'
    assert deliverable.extract_json(reply)["title"] == 'she said "go"'


def test_a_reply_with_no_json_is_refused():
    with pytest.raises(deliverable.PlanInvalid):
        deliverable.extract_json("I could not do that.")


def test_an_empty_reply_is_refused():
    with pytest.raises(deliverable.PlanInvalid):
        deliverable.extract_json("")


# ── Validation ────────────────────────────────────────────────────────────────

def good_plan(n=2):
    return {"title": "Cascade 2.0", "subtitle": "The thesis",
            "slides": [{"heading": f"Slide {i}", "bullets": [f"Point {i}."]}
                       for i in range(n)]}


def test_a_valid_plan_normalizes():
    plan = deliverable.validate_plan(good_plan())
    assert plan["title"] == "Cascade 2.0"
    assert len(plan["slides"]) == 2
    assert plan["slides"][0]["note"] == ""


def test_a_plan_with_no_slides_is_refused():
    with pytest.raises(deliverable.PlanInvalid, match="no slides"):
        deliverable.validate_plan({"title": "T", "slides": []})


def test_a_plan_with_no_title_is_refused():
    with pytest.raises(deliverable.PlanInvalid, match="no title"):
        deliverable.validate_plan({"title": "  ", "slides": [{"heading": "H", "bullets": ["b"]}]})


def test_a_slide_with_no_bullets_is_refused():
    # An empty slide is the shape that gets presented by accident.
    with pytest.raises(deliverable.PlanInvalid, match="no bullets"):
        deliverable.validate_plan({"title": "T", "slides": [{"heading": "H", "bullets": []}]})


def test_a_slide_with_no_heading_is_refused():
    with pytest.raises(deliverable.PlanInvalid, match="no heading"):
        deliverable.validate_plan({"title": "T", "slides": [{"heading": "", "bullets": ["b"]}]})


def test_an_oversized_plan_is_refused_rather_than_truncated():
    # Truncating would lose the last slides silently; the user must see why.
    with pytest.raises(deliverable.PlanInvalid, match="more than"):
        deliverable.validate_plan(good_plan(deliverable.MAX_SLIDES + 1))


def test_a_dash_is_replaced_at_validation():
    # DESIGN.md bars dashes from every generated artifact, and a prompt
    # instruction is never the guarantee.
    plan = deliverable.validate_plan(
        {"title": "A" + chr(0x2014) + "B",
         "slides": [{"heading": "H" + chr(0x2013) + "I", "bullets": ["x" + chr(0x2014) + "y"]}]})
    # json.dumps escapes non-ASCII to \u2014 by default, so a scan of its output
    # can never see the dash: assert over the real strings instead.
    strings = [plan["title"]] + [plan["slides"][0]["heading"]] + plan["slides"][0]["bullets"]
    for value in strings:
        assert chr(0x2014) not in value, value
        assert chr(0x2013) not in value, value
    assert plan["title"] == "A, B"


# ── Building ──────────────────────────────────────────────────────────────────

def test_the_deck_is_native_text_not_a_picture(tmp_path):
    from pptx import Presentation
    out = deliverable.build_pptx(good_plan(), tmp_path / "d.pptx")
    prs = Presentation(str(out))
    # Every string a reader sees is a real text frame they can edit.
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    assert "Cascade 2.0" in texts
    assert any("Point 0." in t for t in texts)
    assert not any(sh.shape_type == 13 for s in prs.slides for sh in s.shapes), \
        "an image means the text was rendered, not written"


def test_the_cover_is_one_slide_beyond_the_plan(tmp_path):
    from pptx import Presentation
    out = deliverable.build_pptx(good_plan(3), tmp_path / "d.pptx")
    assert len(Presentation(str(out)).slides) == 4


def test_a_speaker_note_reaches_the_deck(tmp_path):
    from pptx import Presentation
    plan = {"title": "T", "slides": [{"heading": "H", "bullets": ["b"], "note": "say it plainly"}]}
    out = deliverable.build_pptx(plan, tmp_path / "d.pptx")
    prs = Presentation(str(out))
    assert "say it plainly" in prs.slides[1].notes_slide.notes_text_frame.text


def test_an_invalid_plan_never_reaches_a_file(tmp_path):
    out = tmp_path / "d.pptx"
    with pytest.raises(deliverable.PlanInvalid):
        deliverable.build_pptx({"title": "T", "slides": []}, out)
    assert not out.exists()


def test_a_type_with_no_builder_says_so(tmp_path):
    with pytest.raises(deliverable.PlanInvalid, match="no builder"):
        deliverable.build("xlsx", good_plan(), tmp_path / "d.xlsx")


# ── Verification ──────────────────────────────────────────────────────────────

def test_verification_passes_on_what_the_builder_wrote(tmp_path):
    plan = deliverable.validate_plan(good_plan())
    out = deliverable.build_pptx(plan, tmp_path / "d.pptx")
    detail = deliverable.verify("pptx", plan, out)
    assert "3 slides" in detail


def test_a_missing_file_fails_verification(tmp_path):
    with pytest.raises(deliverable.PlanInvalid, match="no file"):
        deliverable.verify("pptx", good_plan(), tmp_path / "nothing.pptx")


def test_a_slide_count_mismatch_fails_verification(tmp_path):
    built = deliverable.build_pptx(good_plan(2), tmp_path / "d.pptx")
    with pytest.raises(deliverable.PlanInvalid, match="slides"):
        deliverable.verify("pptx", good_plan(5), built)


def test_placeholder_text_fails_verification(tmp_path):
    # A template prompt shipping to a reader is the defect gates never catch.
    plan = {"title": "T", "slides": [{"heading": "H", "bullets": ["TODO: write this"]}]}
    out = deliverable.build_pptx(plan, tmp_path / "d.pptx")
    with pytest.raises(deliverable.PlanInvalid, match="placeholder"):
        deliverable.verify("pptx", plan, out)


def test_verification_is_reached_through_the_real_builder(tmp_path):
    # The verifier must grade the file the builder actually wrote, not a plan
    # projection of it.
    plan = good_plan(4)
    out = deliverable.build("pptx", plan, tmp_path / "d.pptx")
    assert out.exists()
    assert "5 slides" in deliverable.verify("pptx", plan, out)


# ── The stager ────────────────────────────────────────────────────────────────

class FakeRun:
    def __init__(self, stdout="", returncode=0, stderr=""):
        self.stdout, self.returncode, self.stderr = stdout, returncode, stderr


def test_the_stager_returns_a_validated_plan():
    reply = json.dumps(good_plan(3))
    plan = deliverable.stage_plan({"title": "T", "dod": "d"},
                                  run=lambda *a, **k: FakeRun(reply))
    assert len(plan["slides"]) == 3


def test_a_failed_model_call_is_not_a_silent_empty_deck():
    with pytest.raises(deliverable.PlanInvalid, match="model call failed"):
        deliverable.stage_plan({"title": "T"},
                               run=lambda *a, **k: FakeRun("", 1, "quota reached"))


def test_the_prompt_carries_the_context_pack():
    seen = {}

    def run(prompt, **kwargs):
        seen["prompt"] = prompt
        return FakeRun(json.dumps(good_plan()))

    ticket = {"title": "T", "dod": "d", "context_pack": {
        "open_commitments": [{"item": "Send the revised model", "source": "x.md"}],
        "deal_context": [{"text": "Generate wants a term sheet", "source": "h.md"}]}}
    deliverable.stage_plan(ticket, run=run)
    assert "Send the revised model" in seen["prompt"]
    assert "Generate wants a term sheet" in seen["prompt"]


def test_a_revision_comment_outranks_the_rest_of_the_prompt():
    seen = {}

    def run(prompt, **kwargs):
        seen["prompt"] = prompt
        return FakeRun(json.dumps(good_plan()))

    deliverable.stage_plan(
        {"title": "T", "dod": "d", "comments": [{"text": "too long, cut it to five"}]},
        run=run)
    assert "too long, cut it to five" in seen["prompt"]
    assert "take priority" in seen["prompt"]
